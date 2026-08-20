import logging
from decimal import Decimal

from django.db.models import Sum, Count, Q
from django.utils import timezone
from django_filters.rest_framework import DjangoFilterBackend
from rest_framework import generics, status, filters
from rest_framework.permissions import IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework_simplejwt.views import TokenObtainPairView, TokenRefreshView
from apps.api.tasks import parse_payroll_upload

from apps.organizations.models import CheckoffOrganizationMirror, HRUser, AuditLog
from apps.payroll.models import PayrollUpload, SalaryDeduction
from apps.repayments.models import RepaymentBatch, RepaymentRecord
from apps.api.lms_client import get_customer_names_bulk

from .permissions import IsHRUser, IsHRAdmin, BelongsToOrganization
from .serializers import (
    CheckoffOrganizationSerializer,
    HRUserSerializer, HRUserCreateSerializer,
    PayrollUploadSerializer, PayrollUploadCreateSerializer,
    SalaryDeductionSerializer,
    RepaymentBatchSerializer, RepaymentRecordSerializer,
    BatchApproveSerializer, DashboardSummarySerializer,
)
from .throttles import RepaymentRateThrottle, BurstRepaymentThrottle

log = logging.getLogger(__name__)


class HRTokenObtainView(TokenObtainPairView):
    pass


class HRTokenRefreshView(TokenRefreshView):
    pass


class OrgScopedMixin:
    """Mixin that scopes all querysets to request.hr_organization."""
    permission_classes = [IsHRUser]

    def perform_authentication(self, request):
        # JWT auth is lazy — calling super() forces it so request.user is resolved
        # before check_permissions(). We then re-populate hr_user on the underlying
        # Django request so the permission classes can read it via DRF's proxy.
        super().perform_authentication(request)
        request._request.hr_user = None
        request._request.hr_organization = None
        if request.user and request.user.is_authenticated:
            try:
                hr_profile = request.user.hr_profile
                if hr_profile.is_active:
                    request._request.hr_user = hr_profile
                    request._request.hr_organization = hr_profile.organization
            except Exception:
                pass

    def get_org(self):
        if self.request.user.is_superuser:
            code = self.request.query_params.get('org_code')
            if code:
                return CheckoffOrganizationMirror.objects.filter(code=code).first()
        return self.request.hr_organization

    def _audit(self, action, obj_id='', description='', metadata=None):
        AuditLog.objects.create(
            actor=self.request.user,
            organization=self.request.hr_organization,
            action=action,
            object_id=str(obj_id),
            description=description,
            ip_address=self._get_ip(),
            metadata=metadata or {},
        )

    def _get_ip(self):
        xff = self.request.META.get('HTTP_X_FORWARDED_FOR')
        return xff.split(',')[0].strip() if xff else self.request.META.get('REMOTE_ADDR', '')


class DashboardView(OrgScopedMixin, APIView):
    def get(self, request):
        org = self.get_org()
        if not org:
            return Response({'detail': 'No organization found.'}, status=404)

        period = request.query_params.get('period', timezone.now().strftime('%Y-%m'))
        uploads = PayrollUpload.objects.filter(organization=org)
        deductions = SalaryDeduction.objects.filter(organization=org)

        try:
            year, month = map(int, period.split('-'))
            uploads = uploads.filter(payroll_period__year=year, payroll_period__month=month)
            deductions = deductions.filter(deduction_date__year=year, deduction_date__month=month)
        except Exception:
            pass

        agg = deductions.aggregate(
            total_amount=Sum('amount'),
            failed=Count('id', filter=Q(status=SalaryDeduction.STATUS_FAILED)),
        )
        succ_agg = RepaymentRecord.objects.filter(
            organization=org, status=RepaymentRecord.STATUS_SUCCESS,
        ).aggregate(total=Sum('amount_sent'))
        pending_batches = RepaymentBatch.objects.filter(
            organization=org, status=RepaymentBatch.STATUS_DRAFT
        ).count()

        data = {
            'organization': org.name,
            'period': period,
            'total_uploads': uploads.count(),
            'total_deductions': deductions.count(),
            'total_amount': agg['total_amount'] or Decimal('0'),
            'successful_amount': succ_agg['total'] or Decimal('0'),
            'pending_batches': pending_batches,
            'failed_deductions': agg['failed'] or 0,
        }
        return Response(DashboardSummarySerializer(data).data)


class OrganizationDetailView(OrgScopedMixin, generics.RetrieveAPIView):
    serializer_class = CheckoffOrganizationSerializer

    def get_object(self):
        org = self.get_org()
        if not org:
            from rest_framework.exceptions import NotFound
            raise NotFound('Organization not found.')
        return org


class HRUserListCreateView(OrgScopedMixin, generics.ListCreateAPIView):
    def get_permissions(self):
        return [IsHRAdmin()] if self.request.method == 'POST' else [IsHRUser()]

    def get_serializer_class(self):
        return HRUserCreateSerializer if self.request.method == 'POST' else HRUserSerializer

    def get_queryset(self):
        return HRUser.objects.filter(organization=self.get_org()).select_related('user', 'organization')

    def perform_create(self, serializer):
        hr_user = serializer.save()
        self._audit('user_create', hr_user.id, f'Created HR user {hr_user.user.email}')


class HRUserDetailView(OrgScopedMixin, generics.RetrieveUpdateDestroyAPIView):
    permission_classes = [IsHRAdmin]
    serializer_class = HRUserSerializer

    def get_queryset(self):
        return HRUser.objects.filter(organization=self.get_org())

    def perform_destroy(self, instance):
        instance.is_active = False
        instance.user.is_active = False
        instance.save(update_fields=['is_active'])
        instance.user.save(update_fields=['is_active'])
        self._audit('user_deactivate', instance.id, f'Deactivated {instance.user.email}')


class PayrollUploadListCreateView(OrgScopedMixin, generics.ListCreateAPIView):
    parser_classes = [MultiPartParser, FormParser, JSONParser]
    filter_backends = [DjangoFilterBackend, filters.OrderingFilter]
    filterset_fields = ['status', 'payroll_period']
    ordering_fields = ['date_created', 'payroll_period']
    ordering = ['-date_created']

    def get_permissions(self):
        return [IsHRAdmin()] if self.request.method == 'POST' else [IsHRUser()]

    def get_serializer_class(self):
        return PayrollUploadCreateSerializer if self.request.method == 'POST' else PayrollUploadSerializer

    def get_queryset(self):
        return PayrollUpload.objects.filter(organization=self.get_org()).select_related(
            'organization', 'uploaded_by'
        )

    def perform_create(self, serializer):
        upload = serializer.save()
        self._audit('upload', upload.id,
                    f'Uploaded {upload.original_filename} for {upload.payroll_period:%Y-%m}',
                    metadata={'filename': upload.original_filename})

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data, context={'request': request})
        serializer.is_valid(raise_exception=True)
        self.perform_create(serializer)
        return Response(
            PayrollUploadSerializer(serializer.instance, context={'request': request}).data,
            status=status.HTTP_202_ACCEPTED
        )


class PayrollUploadDetailView(OrgScopedMixin, generics.RetrieveDestroyAPIView):
    serializer_class = PayrollUploadSerializer
    permission_classes = [IsHRUser, BelongsToOrganization]

    def get_queryset(self):
        return PayrollUpload.objects.filter(organization=self.get_org())

    def get_permissions(self):
        if self.request.method == 'DELETE':
            return [IsHRAdmin(), BelongsToOrganization()]
        return [IsHRUser(), BelongsToOrganization()]

    def perform_destroy(self, instance):
        if instance.status == PayrollUpload.STATUS_PROCESSING:
            from rest_framework.exceptions import ValidationError
            raise ValidationError(
                'Cannot delete an upload that is currently being processed.'
            )
        file_name = instance.original_filename
        instance.file.delete(save=False)
        instance.delete()
        self._audit(
            'upload_delete',
            instance.id,
            f'Deleted upload {file_name} for {instance.payroll_period:%Y-%m}',
        )


class SalaryDeductionListView(OrgScopedMixin, generics.ListAPIView):
    """
    Full customer names are resolved live from the LMS's exclusive-details
    endpoint (see lms_client.get_customer_names_bulk) rather than a local
    customer table. To keep this cheap, the lookup happens AFTER pagination
    (only for the current page's rows) and only for distinct phone numbers,
    with each result cached in Redis by lms_client.
    """
    serializer_class = SalaryDeductionSerializer
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = ['status', 'deduction_date', 'upload']
    search_fields = ['phone_number', 'employee_name', 'employee_id', 'reference']
    ordering_fields = ['deduction_date', 'amount', 'date_created']
    ordering = ['-date_created']

    def get_queryset(self):
        return SalaryDeduction.objects.filter(organization=self.get_org()).select_related(
            'organization', 'upload'
        )

    def list(self, request, *args, **kwargs):
        queryset = self.filter_queryset(self.get_queryset())
        page = self.paginate_queryset(queryset)
        target = page if page is not None else queryset

        phone_numbers = list({d.phone_number for d in target})
        self._customer_names = get_customer_names_bulk(phone_numbers)

        serializer = self.get_serializer(target, many=True)
        if page is not None:
            return self.get_paginated_response(serializer.data)
        return Response(serializer.data)

    def get_serializer_context(self):
        context = super().get_serializer_context()
        context['customer_names'] = getattr(self, '_customer_names', {})
        return context


class SalaryDeductionDetailView(OrgScopedMixin, generics.RetrieveAPIView):
    serializer_class = SalaryDeductionSerializer
    permission_classes = [IsHRUser, BelongsToOrganization]

    def get_queryset(self):
        return SalaryDeduction.objects.filter(organization=self.get_org())

    def get_serializer_context(self):
        context = super().get_serializer_context()
        # Single-object retrieve — resolve just this one phone number.
        obj = getattr(self, '_object_for_context', None)
        if obj is None:
            try:
                obj = self.get_object()
                self._object_for_context = obj
            except Exception:
                obj = None
        if obj is not None:
            context['customer_names'] = get_customer_names_bulk([obj.phone_number])
        return context


class RepaymentBatchListView(OrgScopedMixin, generics.ListAPIView):
    serializer_class = RepaymentBatchSerializer
    filter_backends = [DjangoFilterBackend, filters.OrderingFilter]
    filterset_fields = ['status']
    ordering_fields = ['date_created', 'total_amount']
    ordering = ['-date_created']

    def get_queryset(self):
        return RepaymentBatch.objects.filter(organization=self.get_org()).select_related(
            'organization', 'approved_by', 'upload'
        )


class RepaymentBatchDetailView(OrgScopedMixin, generics.RetrieveAPIView):
    serializer_class = RepaymentBatchSerializer
    permission_classes = [IsHRUser, BelongsToOrganization]

    def get_queryset(self):
        return RepaymentBatch.objects.filter(organization=self.get_org())


class RepaymentBatchApproveView(OrgScopedMixin, APIView):
    """
    POST /api/v1/batches/{id}/approve/
    HR Admin approves a DRAFT batch — triggers Celery fan-out.
    """
    permission_classes = [IsHRAdmin]
    throttle_classes = [RepaymentRateThrottle, BurstRepaymentThrottle]

    def post(self, request, pk):
        from apps.repayments.tasks import dispatch_repayment_batch

        org = self.get_org()
        try:
            batch = RepaymentBatch.objects.get(id=pk, organization=org)
        except RepaymentBatch.DoesNotExist:
            return Response({'detail': 'Batch not found.'}, status=404)

        if batch.status != RepaymentBatch.STATUS_DRAFT:
            return Response(
                {'detail': f'Batch is "{batch.status}" — only DRAFT batches can be approved.'},
                status=400
            )

        serializer = BatchApproveSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        batch.approve(request.user)
        task = dispatch_repayment_batch.delay(str(batch.id))
        batch.celery_task_id = task.id
        batch.status = RepaymentBatch.STATUS_DISPATCHING
        batch.save(update_fields=['celery_task_id', 'status'])

        self._audit(
            'batch_approve', batch.id,
            f'Approved batch {str(batch.id)[:8]} — {batch.total_deductions} deductions, KES {batch.total_amount}',
            metadata={'task_id': task.id}
        )

        return Response(RepaymentBatchSerializer(batch).data)


class RepaymentRecordListView(OrgScopedMixin, generics.ListAPIView):
    serializer_class = RepaymentRecordSerializer
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = ['status', 'batch']
    search_fields = ['phone_number', 'lms_response_code']
    ordering_fields = ['dispatched_at', 'amount_sent']
    ordering = ['-dispatched_at']

    def get_queryset(self):
        return RepaymentRecord.objects.filter(organization=self.get_org()).select_related(
            'organization', 'batch', 'deduction'
        )


class UploadApprovalCallbackView(OrgScopedMixin, APIView):
    """Callback endpoint for LMS to confirm upload approval (not yet implemented)."""
    permission_classes = [IsHRAdmin]

    def post(self, request, pk):
        return Response({'detail': 'Not implemented.'}, status=status.HTTP_501_NOT_IMPLEMENTED)


class PayrollTemplateDownloadView(APIView):
    """GET /api/v1/uploads/template/ — download a blank payroll upload template."""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from django.http import HttpResponse

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Payroll Upload"

        # ── header row ──────────────────────────────────────────
        headers = ["phone_number", "amount", "deduction_date"]
        ws.append(headers)

        header_fill = PatternFill(start_color="217346", end_color="217346", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        thin = Side(style="thin", color="AAAAAA")
        cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for col, cell in enumerate(ws[1], start=1):
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = cell_border

        # ── example rows ────────────────────────────────────────
        examples = [
            ["254712486391", 4500.00, "10/2/2026"],
            ["254722813047", 7200.00, "10/2/2026"],
            ["254733651829", 3100.00, "10/2/2026"],
        ]
        note_fill = PatternFill(start_color="F0F7F3", end_color="F0F7F3", fill_type="solid")
        for row_data in examples:
            ws.append(row_data)
            for cell in ws[ws.max_row]:
                cell.border = cell_border
                cell.fill = note_fill

        # ── column widths ────────────────────────────────────────
        ws.column_dimensions["A"].width = 22   # phone_number
        ws.column_dimensions["B"].width = 16   # amount
        ws.column_dimensions["C"].width = 18   # deduction_date

        # ── freeze header ────────────────────────────────────────
        ws.freeze_panes = "A2"

        # ── notes sheet ─────────────────────────────────────────
        notes = wb.create_sheet("Instructions")
        instructions = [
            ["Field", "Format", "Example"],
            ["phone_number", "Must start with 254 (no + or spaces)", "254712486391"],
            ["amount",       "Numeric, no currency symbol",           "4500.00"],
            ["deduction_date", "DD/MM/YYYY",                         "10/2/2026"],
        ]
        hdr_fill = PatternFill(start_color="217346", end_color="217346", fill_type="solid")
        hdr_font = Font(color="FFFFFF", bold=True)
        for r_idx, row_data in enumerate(instructions, start=1):
            notes.append(row_data)
            for cell in notes[r_idx]:
                cell.border = cell_border
                if r_idx == 1:
                    cell.fill = hdr_fill
                    cell.font = hdr_font

        notes.column_dimensions["A"].width = 20
        notes.column_dimensions["B"].width = 42
        notes.column_dimensions["C"].width = 18

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="payroll_upload_template.xlsx"'
        wb.save(response)
        return response


class HealthCheckView(APIView):
    permission_classes = []
    authentication_classes = []

    def get(self, request):
        from django.db import connection
        from django.core.cache import cache
        checks = {}
        try:
            connection.ensure_connection()
            checks['database'] = 'ok'
        except Exception as e:
            checks['database'] = str(e)
        try:
            cache.set('health_check', '1', 5)
            checks['cache'] = 'ok'
        except Exception as e:
            checks['cache'] = str(e)
        all_ok = all(v == 'ok' for v in checks.values())
        return Response(
            {'status': 'healthy' if all_ok else 'degraded', 'checks': checks},
            status=200 if all_ok else 503
        )
